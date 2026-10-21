"""Step 3: generate the recruiters' rating forms (Excel, in French).

Form A (blind): every profile-offer pair of the evaluation set, rated 0-3.
Nothing on it says which offers the model chose, and the order of profiles
and offers is shuffled differently for each recruiter, so neither the model's
choices nor a fixed order can steer the ratings.

Form B (sent only once form A has come back): the model's recommendations for
each profile, with its explanation, judged for accuracy and usefulness.
Sending it earlier would show the recruiter the model's choices before the
blind ratings are done.

Usage: python make_forms.py R1 R2 R3      (one pair of files per recruiter)
Output: forms/Evaluation_A_<R>.xlsx and forms/Evaluation_B_<R>.xlsx
"""
from __future__ import annotations

import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from common import DATA, FORMS, read_json

GRADES = ["0", "1", "2", "3"]
ACCURACY = ["Exacte", "Partiellement exacte", "Inexacte"]
USEFUL = ["Oui", "Non"]

REGIONS = {"ADAMAOUA": "Adamaoua", "CENTRE": "Centre", "EST": "Est", "EXTREME_NORD": "Extrême-Nord",
           "LITTORAL": "Littoral", "NORD": "Nord", "NORD_OUEST": "Nord-Ouest", "OUEST": "Ouest",
           "SUD": "Sud", "SUD_OUEST": "Sud-Ouest"}
INDUSTRIES = {"AGRICULTURE": "Agriculture", "BANKING_FINANCE": "Banque et finance",
              "EDUCATION": "Éducation", "HEALTHCARE": "Santé",
              "LOGISTICS_TRANSPORT": "Logistique et transport",
              "MANUFACTURING": "Industrie manufacturière", "TELECOMMUNICATIONS": "Télécommunications"}
TYPES = {"FULL_TIME": "Temps plein", "PART_TIME": "Temps partiel", "CONTRACT": "Contrat à durée déterminée",
         "INTERN": "Stage"}
SITES = {"ONSITE": "Sur site", "HYBRID": "Hybride", "REMOTE": "À distance"}
JOB_LANGUAGES = {"FRENCH": "Français", "ENGLISH": "Anglais", "BILINGUAL": "Bilingue (français et anglais)"}
SPOKEN = {"French": "Français", "English": "Anglais", "Pidgin": "Pidgin"}

HEADER_FILL = PatternFill("solid", fgColor="0D6E5B")
BAND_FILL = PatternFill("solid", fgColor="EEF7F4")
INPUT_FILL = PatternFill("solid", fgColor="FFF6D6")
THIN = Side(style="thin", color="C9D6D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def location(o: dict) -> str:
    parts = [o["city"], REGIONS.get(o["region"], o["region"])]
    return ", ".join(p for p in parts if p)


def salary(o: dict) -> str:
    lo, hi = o["salaryMin"], o["salaryMax"]
    cur = o["salaryCurrency"] or "XAF"
    fmt = lambda v: f"{float(v):,.0f}".replace(",", " ")
    if lo and hi:
        return f"{fmt(lo)} – {fmt(hi)} {cur}" if lo != hi else f"{fmt(lo)} {cur}"
    if lo:
        return f"À partir de {fmt(lo)} {cur}"
    if hi:
        return f"Jusqu'à {fmt(hi)} {cur}"
    return "Non précisé"


def candidate_lines(p: dict) -> str:
    langs = ", ".join(SPOKEN.get(l, l) for l in p["spokenLanguages"])
    return (f"Poste visé : {p['targetRole']}\n"
            f"Compétences : {', '.join(p['skills'])}\n"
            f"Langues : {langs}\n"
            f"Expérience : {p['yearsOfExperience']} an(s)\n"
            f"Région souhaitée : {REGIONS.get(p['preferredRegion'], p['preferredRegion'])}")


def header(ws, titles: list[str], widths: list[int]) -> None:
    ws.append(titles)
    for col, (title, width) in enumerate(zip(titles, widths), 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def instructions(ws, title: str, paragraphs: list[str]) -> None:
    ws.column_dimensions["A"].width = 110
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14, color="0D6E5B")
    for text in paragraphs:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = max(18, 15 * (len(text) // 105 + 1))


def form_a(recruiter: str, evalset: dict) -> Workbook:
    rng = random.Random(f"A-{recruiter}")
    offers = {o["code"]: o for o in evalset["offers"]}
    profiles = list(evalset["profiles"])
    rng.shuffle(profiles)

    wb = Workbook()
    instructions(wb.active, "Évaluation de recommandations d'offres d'emploi — Partie A", [
        "Merci de participer à cette évaluation, réalisée dans le cadre d'un mémoire de Master 2 "
        "(SIGL, Université de Yaoundé I). Elle vise à comparer les recommandations d'offres "
        "produites automatiquement par une plateforme de recrutement avec le jugement de "
        "recruteurs. Durée estimée : 40 minutes.",
        f"Vous êtes l'évaluateur {recruiter}. Vos réponses sont traitées de façon anonyme.",
        "Ce que vous évaluez. L'onglet « Évaluation » présente 100 paires (candidat, offre). "
        "Pour chaque paire, indiquez dans la colonne jaune dans quelle mesure l'offre convient "
        "au candidat, en choisissant une note dans la liste :",
        "0 — Ne convient pas : je ne transmettrais pas cette offre à ce candidat.",
        "1 — Faiblement adaptée : quelques points communs, mais le candidat ne serait "
        "probablement pas retenu.",
        "2 — Adaptée : l'offre correspond au profil ; je la lui recommanderais.",
        "3 — Très adaptée : forte correspondance ; ce serait l'une de mes premières propositions.",
        "Règles. Jugez uniquement à partir des informations affichées (le candidat et l'offre). "
        "Les candidats sont fictifs. Les offres proviennent de la base de démonstration de la "
        "plateforme et ne sont pas des offres réellement publiées par les entreprises citées. "
        "Travaillez seul, sans consulter d'autres évaluateurs, et notez toutes les paires. "
        "La colonne « Commentaire » est facultative.",
        "L'onglet « Candidats » récapitule les dix profils. Une fois terminé, enregistrez le "
        "fichier et renvoyez-le à l'étudiant. Une seconde partie, plus courte, vous sera alors "
        "transmise.",
    ])
    wb.active.title = "Consignes"

    cands = wb.create_sheet("Candidats")
    header(cands, ["Candidat", "Résumé", "Détail du profil"], [11, 55, 70])
    for p in sorted(evalset["profiles"], key=lambda p: p["code"]):
        cands.append([p["code"], p["summaryFr"], candidate_lines(p)])
        for cell in cands[cands.max_row]:
            cell.alignment, cell.border = WRAP, BORDER
        cands.row_dimensions[cands.max_row].height = 80

    ws = wb.create_sheet("Évaluation")
    header(ws, ["N°", "Candidat", "Profil du candidat", "Offre", "Intitulé du poste", "Entreprise",
                "Secteur", "Lieu", "Contrat", "Mode de travail", "Langue de travail", "Salaire",
                "Description", "Note (0 à 3)", "Commentaire (facultatif)"],
           [5, 10, 42, 8, 30, 22, 18, 18, 16, 14, 18, 20, 45, 12, 28])
    grade = DataValidation(type="list", formula1='"' + ",".join(GRADES) + '"', allow_blank=True,
                           showErrorMessage=True, errorTitle="Note invalide",
                           error="Choisissez une note entre 0 et 3.")
    ws.add_data_validation(grade)

    n = 0
    for block, p in enumerate(profiles):
        codes = list(evalset["pools"][p["code"]])
        rng.shuffle(codes)
        for c in codes:
            o = offers[c]
            n += 1
            ws.append([n, p["code"], candidate_lines(p), c, o["title"], o["company"],
                       INDUSTRIES.get(o["industry"], o["industry"] or "Non précisé"), location(o),
                       TYPES.get(o["type"], o["type"]), SITES.get(o["site"], o["site"]),
                       JOB_LANGUAGES.get(o["language"], "Non précisée"), salary(o),
                       o["description"], None, None])
            row = ws.max_row
            for cell in ws[row]:
                cell.alignment, cell.border = WRAP, BORDER
                if block % 2:
                    cell.fill = BAND_FILL
            ws.cell(row=row, column=14).fill = INPUT_FILL
            ws.cell(row=row, column=14).alignment = Alignment(horizontal="center", vertical="top")
            grade.add(ws.cell(row=row, column=14))
            ws.row_dimensions[row].height = 78
    return wb


def form_b(recruiter: str, evalset: dict, runs: dict) -> Workbook:
    rng = random.Random(f"B-{recruiter}")
    offers = {o["code"]: o for o in evalset["offers"]}
    profiles = list(evalset["profiles"])
    rng.shuffle(profiles)

    wb = Workbook()
    instructions(wb.active, "Évaluation de recommandations d'offres d'emploi — Partie B", [
        "Cette seconde partie porte sur les explications fournies par l'intelligence artificielle "
        "de la plateforme. Pour chaque candidat, l'onglet « Explications » présente les offres que "
        "le modèle a recommandées, dans son ordre, avec la phrase qui justifie chaque "
        "recommandation. Durée estimée : 15 minutes.",
        "Pour chaque ligne, indiquez :",
        "Exactitude — « Exacte » si l'explication est vraie au regard du profil et de l'offre ; "
        "« Partiellement exacte » si elle contient une part d'erreur ou d'exagération ; "
        "« Inexacte » si elle affirme quelque chose de faux (une compétence absente, un lieu "
        "erroné, etc.).",
        "Utile — « Oui » si cette explication aiderait le candidat à décider s'il postule.",
        "Comme pour la partie A, jugez uniquement à partir des informations affichées et "
        "renvoyez le fichier à l'étudiant une fois terminé.",
    ])
    wb.active.title = "Consignes"

    ws = wb.create_sheet("Explications")
    header(ws, ["N°", "Candidat", "Profil du candidat", "Rang", "Offre", "Intitulé du poste",
                "Entreprise", "Lieu", "Explication donnée par l'IA", "Exactitude", "Utile ?",
                "Commentaire (facultatif)"],
           [5, 10, 42, 7, 8, 30, 22, 18, 55, 20, 10, 28])
    acc = DataValidation(type="list", formula1='"' + ",".join(ACCURACY) + '"', allow_blank=True)
    use = DataValidation(type="list", formula1='"' + ",".join(USEFUL) + '"', allow_blank=True)
    ws.add_data_validation(acc)
    ws.add_data_validation(use)

    n = 0
    for block, p in enumerate(profiles):
        first_run = runs[p["code"]][0]          # the answer a seeker would have seen
        for rank, pick in enumerate(first_run["picks"], 1):
            o = offers[pick["code"]]
            n += 1
            ws.append([n, p["code"], candidate_lines(p), rank, o["code"], o["title"], o["company"],
                       location(o), pick["reason"], None, None, None])
            row = ws.max_row
            for cell in ws[row]:
                cell.alignment, cell.border = WRAP, BORDER
                if block % 2:
                    cell.fill = BAND_FILL
            for col, dv in ((10, acc), (11, use)):
                ws.cell(row=row, column=col).fill = INPUT_FILL
                dv.add(ws.cell(row=row, column=col))
            ws.row_dimensions[row].height = 78
    return wb


def main() -> None:
    recruiters = sys.argv[1:] or ["R1", "R2", "R3"]
    evalset = read_json(DATA / "evaluation_set.json")
    runs_path = DATA / "gemini_runs.json"
    runs = read_json(runs_path)["runs"] if runs_path.exists() else None

    FORMS.mkdir(parents=True, exist_ok=True)
    for r in recruiters:
        form_a(r, evalset).save(FORMS / f"Evaluation_A_{r}.xlsx")
        print(f"  forms/Evaluation_A_{r}.xlsx")
        if runs and all(p["code"] in runs for p in evalset["profiles"]):
            form_b(r, evalset, runs).save(FORMS / f"Evaluation_B_{r}.xlsx")
            print(f"  forms/Evaluation_B_{r}.xlsx")
        else:
            print("  (form B needs data/gemini_runs.json for every profile: run run_gemini.py first)")


if __name__ == "__main__":
    main()
