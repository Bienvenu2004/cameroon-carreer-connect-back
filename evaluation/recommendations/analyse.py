"""Step 4: compare the model's recommendations with the recruiters' judgements.

Reads the completed forms from responses/ (files named ..._A_<R>.xlsx and
..._B_<R>.xlsx) and writes results/results.md and results/results_table.tex.

Ground truth. A profile-offer pair's relevance is the mean of the recruiters'
grades (0-3); the pair counts as relevant when that mean is at least 2
("suitable, I would recommend it").

Rankers compared, each choosing its top five from the same pool of ten:
  gemini   the platform's recommendations (every run, then averaged)
  keyword  a transparent non-AI baseline: shared words between the candidate's
           skills and target role and the offer's title, description and
           sector, plus half a point each for region and language fit
  random   the expectation of choosing five offers at random

Metrics, per profile then averaged: precision of the recommended offers,
P@5, recall@5 and nDCG@5 (graded gains 2^g - 1). Agreement between
recruiters is quadratic-weighted Cohen's kappa. Gemini and the keyword
baseline are compared with an exact two-sided Wilcoxon signed-rank test on
the per-profile nDCG@5. Stability is the mean Jaccard overlap of the top-five
sets across runs.

--self-test fills the forms with SIMULATED grades to exercise every code path.
Its output goes to the console only and must never be reported as a result.
"""
from __future__ import annotations

import itertools
import math
import random
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from common import DATA, FORMS, RESPONSES, RESULTS, read_json

K = 5
RELEVANT = 2.0
RANDOM_TRIALS = 20000
STOP = {"and", "the", "for", "with", "des", "les", "and", "from", "other"}


# ----------------------------------------------------------------- reading

def recruiter_of(path: Path) -> str:
    m = re.search(r"_(R\d+)", path.stem)
    if not m:
        raise SystemExit(f"cannot tell which recruiter filled {path.name} (expected _R1, _R2...)")
    return m.group(1)


def read_part_a(path: Path) -> dict[tuple[str, str], int]:
    ws = load_workbook(path, data_only=True)["Évaluation"]
    grades = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        profile, offer, grade = row[1], row[3], row[13]
        if not profile or not offer:
            continue
        if grade is None or str(grade).strip() == "":
            continue
        value = int(float(str(grade).strip()))
        if value not in (0, 1, 2, 3):
            raise SystemExit(f"{path.name}: grade {grade!r} for {profile}/{offer} is not 0-3")
        grades[(profile, offer)] = value
    return grades


def read_part_b(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True)["Explications"]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        out.append({"profile": row[1], "rank": row[3], "offer": row[4],
                    "accuracy": (row[9] or "").strip(), "useful": (row[10] or "").strip()})
    return out


# ----------------------------------------------------------------- metrics

def weighted_kappa(a: list[int], b: list[int], categories: int = 4) -> float:
    n = len(a)
    observed = [[0] * categories for _ in range(categories)]
    for x, y in zip(a, b):
        observed[x][y] += 1
    rows = [sum(r) for r in observed]
    cols = [sum(observed[i][j] for i in range(categories)) for j in range(categories)]
    num = den = 0.0
    for i in range(categories):
        for j in range(categories):
            w = (i - j) ** 2 / (categories - 1) ** 2
            num += w * observed[i][j]
            den += w * rows[i] * cols[j] / n
    return 1.0 - num / den if den else float("nan")


def dcg(gains: list[float]) -> float:
    return sum((2 ** g - 1) / math.log2(i + 2) for i, g in enumerate(gains))


def score_ranking(picks: list[str], pool: list[str], truth: dict[str, float]) -> dict:
    relevant = {c for c in pool if truth[c] >= RELEVANT}
    top = picks[:K]
    hits = sum(c in relevant for c in top)
    ideal = dcg(sorted((truth[c] for c in pool), reverse=True)[:K])
    return {
        "precision": hits / len(top) if top else float("nan"),
        "p_at_k": hits / K,
        "recall": hits / min(K, len(relevant)) if relevant else float("nan"),
        "ndcg": dcg([truth[c] for c in top]) / ideal if ideal else float("nan"),
        "returned": len(top),
    }


def tokens(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-zà-ÿ0-9]+", text.lower()) if len(w) >= 3 and w not in STOP}


def keyword_ranking(profile: dict, pool: list[dict]) -> list[str]:
    wanted = tokens(" ".join(profile["skills"] + [profile["targetRole"]]))
    langs = set(profile["spokenLanguages"])

    def score(o: dict) -> float:
        have = tokens(" ".join([o["title"], o["description"], o["industry"].replace("_", " ")]))
        # light stemming: "engineer" matches "engineering"
        shared = sum(any(h[:5] == w[:5] for h in have) for w in wanted)
        region = 0.5 if o["region"] == profile["preferredRegion"] else 0.0
        lang_ok = {"": True, "FRENCH": "French" in langs, "ENGLISH": "English" in langs,
                   "BILINGUAL": {"French", "English"} <= langs}.get(o["language"], True)
        return shared + region + (0.5 if lang_ok else 0.0)

    order = {o["code"]: i for i, o in enumerate(pool)}
    return [o["code"] for o in sorted(pool, key=lambda o: (-score(o), order[o["code"]]))][:K]


def random_expectation(pool: list[str], truth: dict[str, float], rng: random.Random) -> dict:
    totals = defaultdict(float)
    for _ in range(RANDOM_TRIALS):
        picks = rng.sample(pool, K)
        for key, value in score_ranking(picks, pool, truth).items():
            if not math.isnan(value):
                totals[key] += value
    return {k: v / RANDOM_TRIALS for k, v in totals.items()}


def wilcoxon_exact(x: list[float], y: list[float]) -> tuple[float, int]:
    """Two-sided exact Wilcoxon signed-rank p-value; zero differences dropped."""
    diffs = [a - b for a, b in zip(x, y) if abs(a - b) > 1e-12]
    n = len(diffs)
    if n == 0:
        return 1.0, 0
    ordered = sorted(range(n), key=lambda i: abs(diffs[i]))
    ranks = [0.0] * n
    i = 0
    while i < n:
        j = i
        while j + 1 < n and abs(abs(diffs[ordered[j + 1]]) - abs(diffs[ordered[i]])) < 1e-12:
            j += 1
        for k in range(i, j + 1):
            ranks[ordered[k]] = (i + j) / 2 + 1
        i = j + 1
    w_plus = sum(r for r, d in zip(ranks, diffs) if d > 0)
    total = sum(ranks)
    extreme = min(w_plus, total - w_plus)
    count = sum(1 for signs in itertools.product((0, 1), repeat=n)
                if min(s := sum(r for r, g in zip(ranks, signs) if g), total - s) <= extreme + 1e-12)
    return min(1.0, count / 2 ** n), n


def jaccard(a: list[str], b: list[str]) -> float:
    sa, sb = set(a), set(b)
    return len(sa & sb) / len(sa | sb) if sa | sb else 1.0


def mean_sd(values: list[float]) -> str:
    values = [v for v in values if not math.isnan(v)]
    if not values:
        return "n/a"
    sd = statistics.stdev(values) if len(values) > 1 else 0.0
    return f"{statistics.mean(values):.2f} ± {sd:.2f}"


# ----------------------------------------------------------------- analysis

def analyse(grades_by_recruiter: dict[str, dict], part_b: dict[str, list], simulated: bool) -> str:
    evalset = read_json(DATA / "evaluation_set.json")
    runs = read_json(DATA / "gemini_runs.json")
    offers = {o["code"]: o for o in evalset["offers"]}
    pairs = [(p["code"], c) for p in evalset["profiles"] for c in evalset["pools"][p["code"]]]
    recruiters = sorted(grades_by_recruiter)
    out: list[str] = []
    say = out.append

    if simulated:
        say("> **SIMULATED DATA — code check only. These numbers are not results.**\n")

    say("## Ratings collected\n")
    for r in recruiters:
        say(f"- {r}: {len(grades_by_recruiter[r])} of {len(pairs)} pairs rated")
    missing = [p for p in pairs if not any(p in grades_by_recruiter[r] for r in recruiters)]
    if missing:
        say(f"- {len(missing)} pairs have no rating at all and are left out")

    truth: dict[str, dict[str, float]] = defaultdict(dict)
    for p, c in pairs:
        values = [grades_by_recruiter[r][(p, c)] for r in recruiters if (p, c) in grades_by_recruiter[r]]
        if values:
            truth[p][c] = statistics.mean(values)

    say("\n## Agreement between recruiters\n")
    if len(recruiters) < 2:
        say("Only one recruiter: agreement cannot be measured.")
    else:
        say("| Pair | Pairs rated by both | Exact agreement | Within one point | Weighted kappa |")
        say("|---|---|---|---|---|")
        kappas = []
        for a, b in itertools.combinations(recruiters, 2):
            common = [p for p in pairs if p in grades_by_recruiter[a] and p in grades_by_recruiter[b]]
            xa = [grades_by_recruiter[a][p] for p in common]
            xb = [grades_by_recruiter[b][p] for p in common]
            kappa = weighted_kappa(xa, xb)
            kappas.append(kappa)
            exact = sum(x == y for x, y in zip(xa, xb)) / len(common)
            near = sum(abs(x - y) <= 1 for x, y in zip(xa, xb)) / len(common)
            say(f"| {a}–{b} | {len(common)} | {exact:.0%} | {near:.0%} | {kappa:.2f} |")
        say(f"\nMean weighted kappa: **{statistics.mean(kappas):.2f}**")

    all_truth = [v for p in truth for v in truth[p].values()]
    relevant_share = sum(v >= RELEVANT for v in all_truth) / len(all_truth)
    say(f"\nRelevant pairs (mean grade ≥ {RELEVANT:g}): {relevant_share:.0%} of {len(all_truth)}")

    rng = random.Random(7)
    per_profile = []
    stability, hallucinated, fallback_runs, total_runs, latencies = [], 0, Counter(), 0, []
    for profile in evalset["profiles"]:
        code = profile["code"]
        pool = [c for c in evalset["pools"][code] if c in truth[code]]
        if len(pool) < K:
            continue
        gem_runs = [[pk["code"] for pk in r["picks"] if pk["code"] in truth[code]]
                    for r in runs["runs"][code]]
        for r in runs["runs"][code]:
            total_runs += 1
            hallucinated += r.get("hallucinated", 0)
            fallback_runs[r.get("modelVersion") or r.get("model") or "none"] += 1
            if r.get("latencyMs"):
                latencies.append(r["latencyMs"])
        gem_scores = [score_ranking(picks, pool, truth[code]) for picks in gem_runs]
        gem = {k: statistics.mean([s[k] for s in gem_scores if not math.isnan(s[k])] or [float("nan")])
               for k in gem_scores[0]}
        kw_list = keyword_ranking(profile, [offers[c] for c in pool])
        kw = score_ranking(kw_list, pool, truth[code])
        rnd = random_expectation(pool, truth[code], rng)

        # The model may return fewer than five offers, and P@5 / nDCG@5 count
        # each missing slot as a miss. Judging every method on as many offers
        # as the model returned for this candidate removes that handicap.
        rel = {c for c in pool if truth[code][c] >= RELEVANT}
        depth_prec = [sum(c in rel for c in kw_list[:len(picks)]) / len(picks)
                      for picks in gem_runs if picks]
        gem["precision_same_depth"] = gem["precision"]
        kw["precision_same_depth"] = statistics.mean(depth_prec) if depth_prec else float("nan")
        rnd["precision_same_depth"] = len(rel) / len(pool)
        gem["returned_mean"] = statistics.mean(len(p) for p in gem_runs)
        stability.append(statistics.mean(jaccard(a, b) for a, b in itertools.combinations(gem_runs, 2))
                         if len(gem_runs) > 1 else 1.0)
        per_profile.append({"code": code, "role": profile["targetRole"],
                            "relevant": sum(truth[code][c] >= RELEVANT for c in pool),
                            "gemini": gem, "keyword": kw, "random": rnd})

    say("\n## Recommendation quality (top five of ten)\n")
    say("| Ranker | Precision of picks | P@5 | Recall@5 | nDCG@5 |")
    say("|---|---|---|---|---|")
    for name in ("gemini", "keyword", "random"):
        row = [mean_sd([p[name].get(m, float("nan")) for p in per_profile])
               for m in ("precision", "p_at_k", "recall", "ndcg")]
        say(f"| {name} | " + " | ".join(row) + " |")

    returned = [p["gemini"]["returned_mean"] for p in per_profile]
    say(f"\nThe model returned {statistics.mean(returned):.1f} offers per candidate on average "
        f"(from {min(returned):.1f} to {max(returned):.1f}); the other methods always return five.")

    say("\n### At the model's own depth\n")
    say("Each method judged on as many offers as the model returned for that candidate, "
        "so that returning fewer offers is neither rewarded nor penalised.\n")
    say("| Ranker | Precision at the model's depth |")
    say("|---|---|")
    for name in ("gemini", "keyword", "random"):
        say(f"| {name} | {mean_sd([p[name]['precision_same_depth'] for p in per_profile])} |")
    pv_depth, n_depth = wilcoxon_exact([p["gemini"]["precision_same_depth"] for p in per_profile],
                                       [p["keyword"]["precision_same_depth"] for p in per_profile])
    say(f"\nGemini vs keyword at the same depth: p = {pv_depth:.3f} "
        f"(exact Wilcoxon, n = {n_depth} candidates with a difference)")

    say("\n### Per profile (nDCG@5)\n")
    say("| Profile | Target role | Relevant in pool | Gemini | Keyword | Random |")
    say("|---|---|---|---|---|---|")
    for p in per_profile:
        say(f"| {p['code']} | {p['role']} | {p['relevant']} | {p['gemini']['ndcg']:.2f} | "
            f"{p['keyword']['ndcg']:.2f} | {p['random']['ndcg']:.2f} |")

    g = [p["gemini"]["ndcg"] for p in per_profile]
    pv_kw, n_kw = wilcoxon_exact(g, [p["keyword"]["ndcg"] for p in per_profile])
    pv_rn, n_rn = wilcoxon_exact(g, [p["random"]["ndcg"] for p in per_profile])
    say("\n### Significance (exact Wilcoxon signed-rank, two-sided, per-profile nDCG@5)\n")
    say(f"- Gemini vs keyword: p = {pv_kw:.3f} (n = {n_kw} profiles with a difference)")
    say(f"- Gemini vs random: p = {pv_rn:.3f} (n = {n_rn})")

    say("\n## Behaviour of the model\n")
    say(f"- Runs: {total_runs}; answering model: "
        + ", ".join(f"{m} ×{c}" for m, c in fallback_runs.most_common()))
    say(f"- Recommendations pointing outside the pool (dropped): {hallucinated}")
    say(f"- Stability across runs (mean Jaccard of top-five sets): {statistics.mean(stability):.2f}")
    if latencies:
        say(f"- Latency: median {statistics.median(latencies) / 1000:.1f} s, "
            f"max {max(latencies) / 1000:.1f} s")

    if part_b:
        say("\n## Explanations (form B)\n")
        acc, use = Counter(), Counter()
        for rows in part_b.values():
            for row in rows:
                if row["accuracy"]:
                    acc[row["accuracy"]] += 1
                if row["useful"]:
                    use[row["useful"]] += 1
        judged = sum(acc.values())
        for label in ("Exacte", "Partiellement exacte", "Inexacte"):
            say(f"- {label}: {acc[label]} ({acc[label] / judged:.0%})" if judged else f"- {label}: 0")
        if sum(use.values()):
            say(f"- Judged useful: {use['Oui']} of {sum(use.values())} ({use['Oui'] / sum(use.values()):.0%})")

    if not simulated:
        RESULTS.mkdir(parents=True, exist_ok=True)
        with (RESULTS / "results_table.tex").open("w", encoding="utf-8") as tex:
            tex.write("% Generated by analyse.py -- do not edit by hand.\n")
            for name, label in (("gemini", "Gemini"), ("keyword", "Keyword baseline"),
                                ("random", "Random choice")):
                cells = []
                for m in ("precision_same_depth", "p_at_k", "recall", "ndcg"):
                    vals = [p[name].get(m, float("nan")) for p in per_profile]
                    vals = [v for v in vals if not math.isnan(v)]
                    cells.append(f"{statistics.mean(vals):.2f}" if vals else "n/a")
                tex.write(f"{label} & " + " & ".join(cells) + " \\\\\n\\hline\n")
    return "\n".join(out)


def simulate() -> tuple[dict, dict]:
    """Fake recruiters who roughly agree with the keyword baseline, plus noise."""
    evalset = read_json(DATA / "evaluation_set.json")
    offers = {o["code"]: o for o in evalset["offers"]}
    rng = random.Random(1)
    grades = {r: {} for r in ("R1", "R2", "R3")}
    for p in evalset["profiles"]:
        pool = [offers[c] for c in evalset["pools"][p["code"]]]
        liked = set(keyword_ranking(p, pool)[:3])
        for o in pool:
            base = 2.5 if o["code"] in liked else 0.7
            for r in grades:
                grades[r][(p["code"], o["code"])] = max(0, min(3, round(base + rng.gauss(0, 0.8))))
    return grades, {}


def main() -> None:
    if "--self-test" in sys.argv:
        grades, part_b = simulate()
        print(analyse(grades, part_b, simulated=True))
        return

    files_a = sorted(RESPONSES.glob("*_A_R*.xlsx"))
    if not files_a:
        raise SystemExit(f"no completed form A in {RESPONSES} (expected files like Evaluation_A_R1.xlsx)")
    grades = {recruiter_of(f): read_part_a(f) for f in files_a}
    part_b = {recruiter_of(f): read_part_b(f) for f in sorted(RESPONSES.glob("*_B_R*.xlsx"))}
    report = analyse(grades, part_b, simulated=False)
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / "results.md").write_text("# Recommendation evaluation: results\n\n" + report + "\n",
                                        encoding="utf-8")
    print(report)
    print(f"\nwritten: {RESULTS / 'results.md'} and {RESULTS / 'results_table.tex'}")


if __name__ == "__main__":
    main()
